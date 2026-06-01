#!/usr/bin/env python3
"""Generate the DATA 615 Unit 3 Assignment Option B essay as a .docx file."""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_LINE_SPACING


def add_title_page(doc):
    for _ in range(6):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Should AI Strive to Be More Ethical Than Humans?")
    r.bold = True
    r.font.size = Pt(16)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = s.add_run("An Analysis of Ethical Decision-Making Using the MyGoodness Game")
    sr.italic = True
    sr.font.size = Pt(13)

    for _ in range(8):
        doc.add_paragraph()

    for line in [
        "DATA 615 9041: AI Ethics (2265)",
        "Unit 3 Assignment – Option B",
        "University of Maryland Global Campus",
        "June 1, 2026",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line).font.size = Pt(12)

    doc.add_page_break()


def h1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    return p


def body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.space_after = Pt(0)
    return p


def ref(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_after = Pt(0)
    return p


doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
for section in doc.sections:
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

add_title_page(doc)

# ---------------------------------------------------------------- Introduction
h1(doc, "Should AI Strive to Be More Ethical Than Humans?")
body(doc,
    "Artificial intelligence is increasingly entrusted with decisions that once "
    "belonged solely to human moral agents, from triaging medical resources to "
    "recommending where philanthropic dollars should flow. This shift forces a "
    "deceptively simple question: when an AI makes a moral choice, should it merely "
    "reproduce what a typical human would do, or should it aspire to do better? This "
    "essay examines that question through the lens of charitable giving. After playing "
    "MyGoodness—an interactive game developed by the MIT Media Lab that asks players "
    "to allocate hypothetical donations—I reflect on how the exercise exposed my own "
    "giving biases, situate those findings within the scholarly debate on machine "
    "ethics, evaluate the issue through utilitarian, deontological, and virtue-ethics "
    "frameworks, and conclude with concrete recommendations for how an AI donation "
    "agent should be designed. My central argument is that, in the narrow but "
    "consequential domain of charitable giving, an AI should strive to be more ethical "
    "than the average human—not by imposing an alien value system, but by correcting "
    "the predictable cognitive biases that cause well-intentioned people to give less "
    "effectively than they otherwise would.")

# ---------------------------------------------- Part 1: Reflection
h1(doc, "Part I: Reflection on Ethical Decision-Making")

h2(doc, "How I Typically Choose Among Charitable Organizations")
body(doc,
    "Before playing MyGoodness, my own giving was driven far more by emotion and "
    "proximity than by evidence. Like many donors, I gravitate toward causes that feel "
    "personally meaningful: a friend’s fundraising page, a local food bank, a disaster "
    "that dominates the news cycle, or an organization whose brand I already trust. The "
    "decision is usually fast and intuitive rather than analytical—what behavioral "
    "researchers would call a “System 1” response. I rarely compared the cost-"
    "effectiveness of one charity against another, and I almost never asked how many "
    "people a given dollar would actually help. In hindsight, my choices reflected "
    "well-documented psychological tendencies: the “identifiable victim effect,” in "
    "which a single named individual moves us more than statistical lives, and a strong "
    "preference for geographically and culturally close beneficiaries. These instincts "
    "are humane, but they are not optimized for impact.")

h2(doc, "How MyGoodness Clarified My Ethical Preferences")
body(doc,
    "MyGoodness gives each player a hypothetical $1,000 and presents ten rounds in "
    "which $100 must be allocated between competing charitable options (MIT Media Lab, "
    "2017). The game is deliberately constructed as a “charity-based trolley problem”: "
    "each choice isolates a single moral variable—the number of beneficiaries, their "
    "geographic location, the type of aid, or the characteristics of the people helped. "
    "A representative dilemma asks whether to provide nutritious meals to ten children "
    "in North America or clean water to twenty-five people in northern Africa. By "
    "stripping away branding and forcing an explicit trade-off, the game made my "
    "implicit weighting scheme visible. I could see, round after round, that I was "
    "quietly assigning more value to nearby and identifiable beneficiaries than to "
    "larger numbers of distant strangers, even when the distant option helped more "
    "people per dollar. MyGoodness, in other words, turned a vague feeling of "
    "“doing good” into a measurable pattern of revealed preferences, and confronted "
    "me with the discomfort that the game’s designers intended players to feel.")

h2(doc, "Insights From the Feedback on My “Goodness” Compared to Others")
body(doc,
    "The game’s feedback compares each player’s allocations both to the principles of "
    "effective giving—the game was built with The Life You Can Save, the nonprofit "
    "founded by philosopher Peter Singer (MIT Media Lab, 2017)—and to the aggregated "
    "choices of everyone else who has played. Two insights stood out. First, my "
    "instinct-driven choices were measurably less effective than the impartial, impact-"
    "maximizing benchmark; I had repeatedly traded reach for emotional resonance. "
    "Second, and more striking, the “average” player exhibited the very same biases I "
    "did. The crowd, like me, tended to favor proximate, identifiable, and emotionally "
    "vivid causes over interventions that quietly save or improve far more lives. The "
    "key lesson is that popularity is not the same as effectiveness: aggregating many "
    "human donors does not cancel out bias—it reproduces it at scale. That realization "
    "is precisely what makes the AI question so pressing, because a system trained to "
    "imitate the “average” donor would inherit the average donor’s blind spots.")

# ---------------------------------------------- Part 2: AI Ethics Essay
h1(doc, "Part II: The AI Ethics Debate")

h2(doc, "Overview of the Debate")
body(doc,
    "Scholars of machine ethics disagree about how good an artificial moral agent "
    "should be. One camp argues that AI should be only as ethical as humans—that is, "
    "aligned with prevailing human values and preferences. Proponents stress democratic "
    "legitimacy, value pluralism, and the danger of paternalism: an AI that overrides "
    "human moral judgment presumes to know better than the people it serves, raising the "
    "hard problem of whose ethics it would encode (Gabriel, 2020). On this view, "
    "human-compatible AI should defer to human preferences and keep humans firmly in "
    "control (Russell, 2019). The opposing camp argues that AI can and should be more "
    "ethical than humans. Human moral cognition is riddled with biases—scope "
    "insensitivity, in-group favoritism, and the identifiable-victim effect—that lead "
    "us to act against our own considered values. Because a well-designed system is not "
    "subject to fatigue, self-interest, or emotional distortion, it could in principle "
    "reason more consistently and impartially than we do (Wallach & Allen, 2009; "
    "Anderson & Anderson, 2011). The disagreement is not merely academic: the Moral "
    "Machine experiment showed that simply harvesting the “average” human moral "
    "judgment yields culturally inconsistent and sometimes troubling rules (Awad et "
    "al., 2018), which warns against treating the crowd as a moral gold standard.")

h2(doc, "What the MyGoodness Data Implies for AI Giving")
body(doc,
    "My MyGoodness results, and the aggregate data the game displays, supply direct "
    "evidence for this debate. The data show that human donors—individually and on "
    "average—systematically under-weight the most effective interventions because of "
    "predictable biases. If we built an AI donation agent by training it on the average "
    "human’s choices, it would faithfully reproduce scope insensitivity and parochial "
    "preference, channeling money toward emotionally salient but comparatively low-"
    "impact causes. It would, in effect, automate our mistakes. The charitable domain "
    "therefore offers an unusually clean case for the “more ethical” position. Unlike "
    "contested moral dilemmas where reasonable people hold incompatible values, "
    "charitable giving has a relatively uncontroversial shared goal—helping others as "
    "much as possible with limited resources—and a robust evidence base, built by "
    "organizations such as GiveWell and The Life You Can Save, for measuring impact per "
    "dollar (Singer, 2009; MacAskill, 2015). When the goal is agreed and the obstacle is "
    "bias rather than genuine value disagreement, an AI that corrects for bias is not "
    "imposing alien values; it is helping donors achieve the impartial benevolence they "
    "already endorse on reflection. The MyGoodness data thus support an AI that gives "
    "more effectively than the typical human, while cautioning that the AI must still "
    "respect the donor’s legitimate, non-biased preferences.")

h2(doc, "Analysis Through Ethical Theories")
body(doc,
    "Three classical frameworks sharpen this conclusion. Utilitarianism, which judges "
    "actions by the welfare they produce, offers the strongest support for a “more "
    "ethical” AI. Peter Singer’s influential argument that distance and nationality are "
    "morally irrelevant to our duty to relieve suffering implies that giving should "
    "maximize impartial impact per dollar—exactly the standard MyGoodness uses (Singer, "
    "1972, 2009). On a utilitarian view, an AI that overcomes scope insensitivity is not "
    "just permissible but obligatory, because it relieves more suffering. Deontology, by "
    "contrast, supplies essential limits. A Kantian respect for persons as ends "
    "demands that the system honor the donor’s autonomy, secure informed consent, and "
    "operate transparently rather than manipulating the giver toward an outcome the "
    "designer prefers. Deontology reminds us that maximizing good results cannot justify "
    "overriding a donor’s sincere, reasoned commitments—for example, to a cause of "
    "personal significance. Virtue ethics adds a third consideration centered on the "
    "moral character of the giver. Aristotle held that we become generous by practicing "
    "generosity and exercising practical wisdom (phronesis); fully outsourcing moral "
    "judgment to a machine risks atrophying the very virtues that giving is meant to "
    "cultivate. The most defensible position is therefore pluralistic: a framework "
    "anchored in utilitarian effectiveness, bounded by deontological respect for "
    "autonomy and transparency, and tempered by a virtue-ethics commitment to "
    "preserving the human donor’s active moral agency.")

# ---------------------------------------------- Part 3: Recommendations
h1(doc, "Part III: Recommendations for AI Charitable Decision-Making")
body(doc,
    "Translating this analysis into design, an AI charitable-giving system should "
    "follow six principles that balance ethical aspiration with practical reality.")

recs = [
    ("Default to evidence-based effectiveness.",
     "The system’s baseline recommendation should maximize impartial impact per "
     "dollar, drawing on transparent, regularly updated impact evaluations. This is how "
     "the AI legitimately becomes “more ethical than the average human”—by correcting "
     "scope insensitivity and proximity bias rather than amplifying them."),
    ("Preserve donor autonomy through configurability.",
     "Donors should be able to set their own values, causes, and constraints (for "
     "example, weighting a region or cause of personal importance). The AI advises and "
     "optimizes within the donor’s stated commitments rather than overriding them, "
     "honoring the deontological duty to respect persons as ends."),
    ("Guarantee transparency and explainability.",
     "Like MyGoodness itself, the system should show its reasoning—who is helped, how "
     "many, at what cost, and what trade-offs were made—so the donor can understand, "
     "challenge, and learn from each recommendation."),
    ("Keep a human in the loop.",
     "Final allocation decisions, especially large or value-laden ones, should remain "
     "with the donor. Preserving human oversight protects against error and sustains the "
     "donor’s moral development, consistent with virtue ethics."),
    ("Avoid bias amplification and value imposition.",
     "Designers must not train the agent simply to mimic average human choices, which "
     "would entrench documented biases. At the same time, the objective function should "
     "reflect broadly shared, openly justified values—not the private preferences of a "
     "single developer—to maintain legitimacy and pluralism."),
    ("Build in accountability and auditing.",
     "Independent auditing of outcomes, data sources, and recommendation patterns "
     "should be standard, with clear lines of responsibility when the system errs, so "
     "that effectiveness is never pursued at the expense of trust."),
]
for i, (head, text) in enumerate(recs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Inches(0.25)
    rb = p.add_run(f"{i}. {head} ")
    rb.bold = True
    p.add_run(text)

# ---------------------------------------------- Conclusion
h1(doc, "Conclusion")
body(doc,
    "MyGoodness reveals an uncomfortable truth: even thoughtful, generous people give "
    "in predictably biased ways, and the average of many givers is no wiser than the "
    "individual. In the specific domain of charitable donations—where the goal of "
    "helping others effectively is widely shared and impact is measurable—an AI should "
    "indeed strive to be more ethical than the average human. The most ethical course is "
    "not for the AI to impose a single conception of the good, but to help donors live "
    "up to the impartial benevolence they already endorse, while respecting their "
    "autonomy, explaining its reasoning, and keeping human judgment in the loop. Built "
    "on a pluralistic foundation—utilitarian in aim, deontological in its limits, and "
    "virtue-minded in preserving human moral agency—such a system would not replace "
    "human goodness but extend it.")

# ---------------------------------------------- References
doc.add_page_break()
h1(doc, "References")
refs = [
    "Anderson, M., & Anderson, S. L. (Eds.). (2011). Machine ethics. Cambridge "
    "University Press.",
    "Awad, E., Dsouza, S., Kim, R., Schulz, J., Henrich, J., Shariff, A., Bonnefon, "
    "J.-F., & Rahwan, I. (2018). The Moral Machine experiment. Nature, 563(7729), "
    "59–64. https://doi.org/10.1038/s41586-018-0637-6",
    "Gabriel, I. (2020). Artificial intelligence, values, and alignment. Minds and "
    "Machines, 30(3), 411–437. https://doi.org/10.1007/s11023-020-09539-2",
    "MacAskill, W. (2015). Doing good better: How effective altruism can help you make "
    "a difference. Avery.",
    "MIT Media Lab. (2017, December 15). MyGoodness: Making charitable giving more "
    "effective. MIT News. https://news.mit.edu/2017/mit-media-lab-mygoodness-game-"
    "teaches-charitable-effectiveness-1215",
    "Russell, S. (2019). Human compatible: Artificial intelligence and the problem of "
    "control. Viking.",
    "Singer, P. (1972). Famine, affluence, and morality. Philosophy & Public Affairs, "
    "1(3), 229–243.",
    "Singer, P. (2009). The life you can save: Acting now to end world poverty. Random "
    "House.",
    "Wallach, W., & Allen, C. (2009). Moral machines: Teaching robots right from wrong. "
    "Oxford University Press.",
]
for r in refs:
    ref(doc, r)

out = "/home/user/Claude-Code/AI_Ethics_MyGoodness_Essay.docx"
doc.save(out)
print("Saved:", out)
